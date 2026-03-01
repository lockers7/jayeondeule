#!/usr/bin/env python3
"""시스템 파일 목록 Excel + PDF 생성 스크립트"""

import os
import fitz  # PyMuPDF
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# ─── 데이터 정의 ───
MAIN_DATA = [
    (1,  "시작",       "/",               "scheduler.py",            "스케줄러 프로세스 진입점",                        "최초실행 (PID관리)"),
    (2,  "시작",       "/",               "startup.py",              "앱 초기화/종료 (Ollama·ChromaDB·PG 연결)",        "initialize_app()"),
    (3,  "시작",       "/",               "__init__.py",             "패키지 공통 export (settings, logger, db)",       ""),
    (4,  "설정",       "config/",         "settings.py",             "환경변수 기반 앱 설정 로드 (싱글톤)",             ".env 연동"),
    (5,  "설정",       "config/",         "constants.py",            "전역 상수·임계값·스케줄링 주기 정의",             ""),
    (6,  "설정",       "config/",         "mappers.py",              "릴레이·센서 필드 매핑 딕셔너리",                  "재배사별 매핑"),
    (7,  "로그",       "/",               "logs.py",                 "통합 로깅 (일별 로테이션·크기 관리)",             "모듈별 로거 생성"),
    (8,  "DB",         "src/postgresql/",  "connection.py",          "PostgreSQL 연결풀 싱글톤 관리",                   "psycopg2"),
    (9,  "DB",         "src/postgresql/",  "queries.py",             "SQL 쿼리 정의 (CREATE·SELECT·INSERT)",            ""),
    (10, "DB",         "src/postgresql/",  "reader.py",              "DB 데이터 조회 (센서·릴레이·생육·설정)",          ""),
    (11, "벡터DB",     "src/chroma/",      "config.py",              "ChromaDB 호스트·포트·컬렉션명 설정",              "port:8000"),
    (12, "벡터DB",     "src/chroma/",      "client.py",              "ChromaDB REST API 통신·컬렉션 자동생성",          "heartbeat 체크"),
    (13, "벡터DB",     "src/chroma/",      "collections.py",         "컬렉션별 접근 함수 (farm·doc·conv·web)",          "4개 컬렉션"),
    (14, "벡터DB",     "src/chroma/",      "operations.py",          "벡터 CRUD (임베딩·검색·업서트·삭제)",             ""),
    (15, "벡터DB",     "src/chroma/",      "utils.py",               "메타데이터 정제·임베딩 차원 관리",                ""),
    (16, "벡터DB",     "src/chroma/",      "loader.py",              "대량 데이터 벡터DB 일괄 로드",                    "배치 로딩"),
    (17, "스케줄",     "src/control/",     "task_scheduler.py",      "APScheduler 기반 백그라운드 작업 관리",           "6개 기본작업 등록"),
    (18, "수동환경제어","src/control/",     "control_common.py",      "제어 공통상수·임계값·핀매핑·포맷팅",              "64케이스 기초"),
    (19, "수동환경제어","src/control/",     "manual_control.py",      "센서 기반 자동환경제어 (온도·습도·CO2)",          "10초 주기"),
    (20, "AI환경제어", "src/control/",     "ai_control.py",           "LLM 기반 종합분석 환경제어·비상모니터링",         "5분 주기"),
    (21, "릴레이관리", "src/control/",     "relay_manager.py",        "릴레이 값 설정·상태 조회·상태 로깅",             "DB 릴레이 갱신"),
    (22, "스케줄제어", "src/control/",     "schedule_control.py",     "조명·관수 시간대별 자동 스케줄 제어",            "5분 주기"),
    (23, "생육RAG",   "src/ai/learning/",  "growth_rag_processor.py","생육 데이터 기반 인과관계 RAG 문서 생성",         "0시·12시 실행"),
    (24, "학습",       "src/ai/learning/", "data_analyzer.py",        "센서 데이터 패턴 분석",                          ""),
    (25, "학습",       "src/ai/learning/", "model_trainer.py",        "AI 모델 재학습 처리",                            ""),
    (26, "AI/LLM",    "src/ai/",           "llm_client.py",           "Ollama LLM 통신·Tool Use·응답필터링",            "핵심 모듈"),
    (27, "AI/LLM",    "src/ai/",           "query_handler_simple.py", "Tool Use 방식 질의처리 (동기·SSE스트림)",        "핵심 모듈"),
    (28, "AI/LLM",    "src/ai/",           "tools_definition.py",     "LLM용 도구(함수) 스키마 정의",                  "Tool Use 정의"),
    (29, "AI/LLM",    "src/ai/",           "tools_executor.py",       "LLM 요청 도구 실행 (웹검색·DB·URL)",            "핵심 모듈"),
    (30, "MCP통신",   "src/ai/",           "mcp_client.py",           "MCP서버 통신 (web-search·postgres·fetch)",       "JSON-RPC"),
    (31, "대화관리",  "src/ai/",           "conversation_store.py",   "PostgreSQL 기반 멀티턴 대화 저장·조회",          "인메모리 폴백"),
    (32, "통계",      "src/ai/",           "stats_collector.py",      "LLM 응답시간·도구사용·검색성공률 수집",          "Thread-safe"),
    (33, "RAG",       "src/ai/rag/",       "embedder.py",             "텍스트→벡터 임베딩 변환 (bge-m3)",              ""),
    (34, "RAG",       "src/ai/rag/",       "chunker.py",              "문서 의미단위 청킹 (한국어 문장분리)",           ""),
    (35, "RAG",       "src/ai/rag/",       "document_processor.py",   "PDF·CSV·Excel 문서 처리 및 벡터DB 저장",        "문서유형 자동감지"),
    (36, "RAG",       "src/ai/rag/",       "reranker.py",             "벡터 검색결과 재정렬·최적화",                    ""),
    (37, "RAG",       "src/ai/rag/",       "document_enricher.py",    "문서 메타데이터 강화·보충",                      ""),
    (38, "파일처리",  "src/ai/",           "file_processor.py",       "업로드 파일 읽기 (CSV·Excel·PDF·텍스트)",        "100MB 제한"),
    (39, "API",       "api/",             "__main__.py",              "Uvicorn으로 FastAPI 서버 실행 진입점",            "port:8002"),
    (40, "API",       "api/",             "app.py",                   "REST API 라우팅·인증·미들웨어·엔드포인트",       "핵심 모듈"),
    (41, "API",       "api/",             "models.py",                "Pydantic 요청/응답 데이터 모델 정의",            ""),
    (42, "API",       "api/",             "example_client.py",        "API 테스트용 예제 클라이언트",                   "개발용"),
    (43, "음성처리",  "api/",             "voice_router.py",          "음성 입출력 엔드포인트 (STT·TTS)",              "FastAPI Router"),
    (44, "음성처리",  "src/voice/",       "stt_engine.py",            "음성→텍스트 변환 (faster-whisper)",              "CPU int8 양자화"),
    (45, "음성처리",  "src/voice/",       "tts_engine.py",            "텍스트→음성 변환",                               ""),
    (46, "유틸리티",  "src/utils/",       "conversion.py",            "단위·데이터타입 변환 (safe_float 등)",            ""),
    (47, "유틸리티",  "src/utils/",       "validators.py",            "센서값 정제·부울 파싱·유효성 검사",              ""),
    (48, "유틸리티",  "src/utils/",       "date_utils.py",            "날짜·시간 처리 유틸리티",                         ""),
]

SCHEDULE_DATA = [
    (1, "수동 환경제어",        "10초",     "control_all_manual()"),
    (2, "AI 환경제어",          "5분",      "control_all_ai()"),
    (3, "생육 RAG 처리",        "0시·12시", "run_growth_rag()"),
    (4, "조명·관수 스케줄 제어", "5분",     "control_all_schedules()"),
    (5, "통계 수집",            "10분",     "StatsCollector.collect()"),
    (6, "로그 정리",            "03:00",    "delete_old_daily_logs()"),
]

API_DATA = [
    (1, "GET",  "/health",              "서버 상태 확인"),
    (2, "GET",  "/stats",               "시스템 통계 조회"),
    (3, "POST", "/api/v1/query",        "LLM 질의 (동기)"),
    (4, "POST", "/api/v1/query/stream", "LLM 질의 (SSE 스트리밍)"),
    (5, "POST", "/api/v1/rag/perform",  "RAG 검색 수행"),
    (6, "POST", "/api/v1/rag/save",     "RAG 데이터 저장"),
    (7, "POST", "/api/v1/voice/stt",    "음성→텍스트 변환"),
    (8, "POST", "/api/v1/voice/tts",    "텍스트→음성 변환"),
]

# 업무구분별 RGB 색상
CAT_COLORS = {
    "시작":       (0.84, 0.89, 0.94),
    "설정":       (0.89, 0.94, 0.85),
    "로그":       (0.89, 0.94, 0.85),
    "DB":         (1.00, 0.95, 0.80),
    "벡터DB":     (1.00, 0.95, 0.80),
    "스케줄":     (0.99, 0.89, 0.84),
    "수동환경제어":(0.99, 0.89, 0.84),
    "AI환경제어": (0.99, 0.89, 0.84),
    "릴레이관리": (0.99, 0.89, 0.84),
    "스케줄제어": (0.99, 0.89, 0.84),
    "생육RAG":    (0.99, 0.89, 0.84),
    "학습":       (0.99, 0.89, 0.84),
    "AI/LLM":     (0.89, 0.85, 0.95),
    "MCP통신":    (0.89, 0.85, 0.95),
    "대화관리":   (0.89, 0.85, 0.95),
    "통계":       (0.89, 0.85, 0.95),
    "RAG":        (0.89, 0.85, 0.95),
    "파일처리":   (0.89, 0.85, 0.95),
    "API":        (0.84, 0.96, 0.89),
    "음성처리":   (0.84, 0.96, 0.89),
    "유틸리티":   (0.95, 0.95, 0.95),
}
CAT_COLORS_HEX = {
    "시작":"D6E4F0","설정":"E2EFDA","로그":"E2EFDA","DB":"FFF2CC","벡터DB":"FFF2CC",
    "스케줄":"FCE4D6","수동환경제어":"FCE4D6","AI환경제어":"FCE4D6","릴레이관리":"FCE4D6",
    "스케줄제어":"FCE4D6","생육RAG":"FCE4D6","학습":"FCE4D6",
    "AI/LLM":"E2D9F3","MCP통신":"E2D9F3","대화관리":"E2D9F3","통계":"E2D9F3",
    "RAG":"E2D9F3","파일처리":"E2D9F3",
    "API":"D5F5E3","음성처리":"D5F5E3","유틸리티":"F2F2F2",
}

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(SCRIPT_DIR, "시스템파일목록.xlsx")
PDF_PATH = os.path.join(SCRIPT_DIR, "시스템파일목록.pdf")
FONT_PATH = os.path.join(SCRIPT_DIR, "NanumGothic-Regular.ttf")
FONT_BOLD_PATH = os.path.join(SCRIPT_DIR, "NanumGothic-Bold.ttf")


# ═══════════════════════════════════════════════════════════════
#  Excel 생성
# ═══════════════════════════════════════════════════════════════
def generate_excel():
    wb = Workbook()
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_font = Font(name="맑은 고딕", bold=True, size=10, color="FFFFFF")
    hdr_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c_font = Font(name="맑은 고딕", size=9)
    c_align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c_align_l = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws = wb.active
    ws.title = "시스템 파일 목록"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.5

    ws.merge_cells("A1:F1")
    ws["A1"].value = "자연들에 스마트팜 시스템 (agri_ai_core) 파일 목록"
    ws["A1"].font = Font(name="맑은 고딕", bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A2:F2")
    ws["A2"].value = "기준경로: agri_ai_core/  |  총 파일 수: 48개  |  작성일: 2026-03-01"
    ws["A2"].font = Font(name="맑은 고딕", bold=True, size=9, color="666666")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    headers = ["순번", "업무구분", "디렉토리", "파일명", "파일기능설명", "비고"]
    widths = [6, 14, 20, 26, 46, 18]
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=4, column=ci, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = hdr_font, hdr_fill, hdr_align, border
        ws.column_dimensions[get_column_letter(ci)].width = w

    for ri, (no, cat, d, f, desc, note) in enumerate(MAIN_DATA, 5):
        vals = [no, cat, d, f, desc, note]
        cc = CAT_COLORS_HEX.get(cat, "FFFFFF")
        fill = PatternFill(start_color=cc, end_color=cc, fill_type="solid")
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.font, cell.border, cell.fill = c_font, border, fill
            cell.alignment = c_align_c if ci <= 2 else c_align_l

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    for r in range(4, 5 + len(MAIN_DATA)):
        ws.row_dimensions[r].height = 18

    # Sheet 2
    ws2 = wb.create_sheet("스케줄러 및 API")
    ws2.page_setup.orientation = "landscape"
    ws2.page_setup.paperSize = ws2.PAPERSIZE_A4

    ws2.merge_cells("A1:D1")
    ws2["A1"].value = "스케줄러 등록 작업 목록"
    ws2["A1"].font = Font(name="맑은 고딕", bold=True, size=12)

    for ci, (h, w) in enumerate(zip(["순번","작업명","주기","실행 함수"], [8,28,14,36]), 1):
        cell = ws2.cell(row=3, column=ci, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = hdr_font, hdr_fill, hdr_align, border
        ws2.column_dimensions[get_column_letter(ci)].width = w
    for ri, (no, name, period, func) in enumerate(SCHEDULE_DATA, 4):
        for ci, v in enumerate([no, name, period, func], 1):
            cell = ws2.cell(row=ri, column=ci, value=v)
            cell.font, cell.border = c_font, border
            cell.alignment = c_align_c if ci in (1,3) else c_align_l

    ar = 4 + len(SCHEDULE_DATA) + 2
    ws2.merge_cells(f"A{ar}:D{ar}")
    ws2[f"A{ar}"].value = "API 엔드포인트 목록"
    ws2[f"A{ar}"].font = Font(name="맑은 고딕", bold=True, size=12)
    hr = ar + 2
    for ci, h in enumerate(["순번","메서드","경로","기능"], 1):
        cell = ws2.cell(row=hr, column=ci, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = hdr_font, hdr_fill, hdr_align, border
    for ri, (no, m, p, desc) in enumerate(API_DATA, hr+1):
        for ci, v in enumerate([no, m, p, desc], 1):
            cell = ws2.cell(row=ri, column=ci, value=v)
            cell.font, cell.border = c_font, border
            cell.alignment = c_align_c if ci in (1,2) else c_align_l
            cc = "E8F5E9" if m == "GET" else "FFF3E0"
            cell.fill = PatternFill(start_color=cc, end_color=cc, fill_type="solid")

    wb.save(EXCEL_PATH)
    print(f"Excel 생성 완료: {EXCEL_PATH}")


# ═══════════════════════════════════════════════════════════════
#  PDF 생성 (PyMuPDF)
# ═══════════════════════════════════════════════════════════════
# A4 landscape: 842 x 595 pt
PW, PH = 842, 595
MARGIN_L, MARGIN_R, MARGIN_T, MARGIN_B = 28, 28, 28, 28
CONTENT_W = PW - MARGIN_L - MARGIN_R

# 메인 테이블 컬럼 비율 및 너비
MAIN_COL_RATIOS = [0.04, 0.10, 0.14, 0.18, 0.40, 0.14]
MAIN_COL_W = [CONTENT_W * r for r in MAIN_COL_RATIOS]

HDR_COLOR = (0.184, 0.329, 0.588)   # #2F5496
HDR_TEXT = (1, 1, 1)
LINE_COLOR = (0.78, 0.78, 0.78)
ZEBRA = (0.97, 0.97, 0.97)


def _draw_rect(page, x, y, w, h, fill_color):
    """배경 사각형"""
    rect = fitz.Rect(x, y, x + w, y + h)
    shape = page.new_shape()
    shape.draw_rect(rect)
    shape.finish(fill=fill_color, color=LINE_COLOR, width=0.4)
    shape.commit()


def _draw_text(page, x, y, w, h, text, font, fontsize, align="left", bold_font=None):
    """셀 안에 텍스트 (수직 중앙정렬)"""
    use_font = bold_font if bold_font else font
    text_w = use_font.text_length(text, fontsize=fontsize)

    pad = 4
    if align == "center":
        tx = x + (w - text_w) / 2
    elif align == "right":
        tx = x + w - text_w - pad
    else:
        tx = x + pad

    ty = y + (h + fontsize * 0.7) / 2  # 수직 중앙
    tw = fitz.TextWriter(page.rect)
    tw.append((tx, ty), text, font=use_font, fontsize=fontsize)
    tw.write_text(page)


def _draw_header_row(page, y, col_widths, headers, font_bold, fontsize=8):
    """헤더 행"""
    x = MARGIN_L
    h = 18
    for w, txt in zip(col_widths, headers):
        _draw_rect(page, x, y, w, h, HDR_COLOR)
        _draw_text(page, x, y, w, h, txt, font_bold, fontsize, align="center", bold_font=font_bold)
        x += w
    return y + h


def _draw_data_row(page, y, col_widths, values, aligns, fill_color, font, fontsize=7.5):
    """데이터 행"""
    x = MARGIN_L
    h = 14.5
    for w, txt, a in zip(col_widths, values, aligns):
        _draw_rect(page, x, y, w, h, fill_color)
        _draw_text(page, x, y, w, h, str(txt), font, fontsize, align=a)
        x += w
    return y + h


def generate_pdf():
    font = fitz.Font(fontfile=FONT_PATH)
    font_bold = fitz.Font(fontfile=FONT_BOLD_PATH)

    doc = fitz.open()

    # ════════════════════════════════════════════
    #  Page 1~2: 메인 파일 목록
    # ════════════════════════════════════════════
    page = doc.new_page(width=PW, height=PH)
    y = MARGIN_T

    # 제목
    tw = fitz.TextWriter(page.rect)
    title = "자연들에 스마트팜 시스템 (agri_ai_core) 파일 목록"
    title_w = font_bold.text_length(title, fontsize=15)
    tw.append(((PW - title_w) / 2, y + 15), title, font=font_bold, fontsize=15)
    tw.write_text(page)
    y += 28

    # 부제
    tw2 = fitz.TextWriter(page.rect)
    sub = "기준경로: agri_ai_core/  |  총 파일 수: 48개  |  작성일: 2026-03-01"
    sub_w = font.text_length(sub, fontsize=8)
    tw2.append(((PW - sub_w) / 2, y + 8), sub, font=font, fontsize=8)
    tw2.write_text(page, color=(0.4, 0.4, 0.4))
    y += 18

    # 헤더
    headers = ["순번", "업무구분", "디렉토리", "파일명", "파일기능설명", "비고"]
    aligns = ["center", "center", "left", "left", "left", "left"]
    y = _draw_header_row(page, y, MAIN_COL_W, headers, font_bold)

    # 데이터 행
    for no, cat, dir_, file, desc, note in MAIN_DATA:
        if y + 16 > PH - MARGIN_B - 14:
            # 페이지 넘김
            _draw_footer(page, doc.page_count, font)
            page = doc.new_page(width=PW, height=PH)
            y = MARGIN_T
            y = _draw_header_row(page, y, MAIN_COL_W, headers, font_bold)

        fill = CAT_COLORS.get(cat, (1, 1, 1))
        vals = [str(no), cat, dir_, file, desc, note]
        y = _draw_data_row(page, y, MAIN_COL_W, vals, aligns, fill, font)

    _draw_footer(page, doc.page_count, font)

    # ════════════════════════════════════════════
    #  Page 3: 스케줄러 + API
    # ════════════════════════════════════════════
    page3 = doc.new_page(width=PW, height=PH)
    y = MARGIN_T

    # 스케줄러 제목
    tw3 = fitz.TextWriter(page3.rect)
    tw3.append((MARGIN_L, y + 13), "스케줄러 등록 작업 목록", font=font_bold, fontsize=13)
    tw3.write_text(page3)
    y += 26

    sched_ratios = [0.06, 0.28, 0.14, 0.35]
    total_r = sum(sched_ratios)
    sched_w = [CONTENT_W * 0.83 * r / total_r for r in sched_ratios]
    sched_h = ["순번", "작업명", "주기", "실행 함수"]
    sched_a = ["center", "left", "center", "left"]

    y = _draw_header_row(page3, y, sched_w, sched_h, font_bold, fontsize=8.5)

    for i, (no, name, period, func) in enumerate(SCHEDULE_DATA):
        fill = ZEBRA if i % 2 == 0 else (1, 1, 1)
        y = _draw_data_row(page3, y, sched_w, [str(no), name, period, func], sched_a, fill, font, fontsize=8.5)

    y += 30

    # API 제목
    tw4 = fitz.TextWriter(page3.rect)
    tw4.append((MARGIN_L, y + 13), "API 엔드포인트 목록", font=font_bold, fontsize=13)
    tw4.write_text(page3)
    y += 26

    api_ratios = [0.06, 0.10, 0.28, 0.39]
    total_r = sum(api_ratios)
    api_w = [CONTENT_W * 0.83 * r / total_r for r in api_ratios]
    api_hdrs = ["순번", "메서드", "경로", "기능"]
    api_a = ["center", "center", "left", "left"]

    y = _draw_header_row(page3, y, api_w, api_hdrs, font_bold, fontsize=8.5)

    for no, method, path, desc in API_DATA:
        fill = (0.91, 0.96, 0.91) if method == "GET" else (1.0, 0.95, 0.88)
        y = _draw_data_row(page3, y, api_w, [str(no), method, path, desc], api_a, fill, font, fontsize=8.5)

    _draw_footer(page3, doc.page_count, font)

    doc.subset_fonts()
    doc.ez_save(PDF_PATH, garbage=4, deflate=True)
    sz = os.path.getsize(PDF_PATH)
    print(f"PDF 생성 완료: {PDF_PATH} ({sz/1024:.0f}KB)")


def _draw_footer(page, page_num, font):
    """페이지 하단 페이지 번호"""
    tw = fitz.TextWriter(page.rect)
    txt = f"- {page_num} -"
    tw_w = font.text_length(txt, fontsize=7)
    tw.append(((PW - tw_w) / 2, PH - 14), txt, font=font, fontsize=7)
    tw.write_text(page, color=(0.6, 0.6, 0.6))


if __name__ == "__main__":
    generate_excel()
    generate_pdf()
